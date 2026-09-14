"""Patch the 28-August presentation: slides 14, 20, 24, 25.

 14 — grammar + technical wording
 20 — proper subscripts on S_i, P_i, P~, D_i
 24 — superscript on W^(t+1), subscripts on W_i / W_N, spacing fixes
 25 — status/future rewritten to match the xMK-CKKS narrative now in the deck

Run:
  "/c/Users/.../Python312/python.exe" compare/patch_presentation.py
"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

DECK = Path(__file__).resolve().parent.parent / "Presentation" / "Presentation 28th August.pptx"

SUB = "-25000"   # OOXML baseline for subscript
SUP = "30000"    # OOXML baseline for superscript


# ---------------------------------------------------------------- helpers
def para_font(para):
    """Capture (size, bold, name) from the first run so we can rebuild in style."""
    for r in para.runs:
        return (r.font.size, r.font.bold, r.font.name)
    return (None, None, None)


def rebuild(para, parts, base=None):
    """Replace a paragraph's runs with `parts`.

    parts: list of (text, style) where style in
           {"", "b", "sub", "sup", "bsub", "bsup"}
    """
    size, bold, name = base if base else para_font(para)

    # drop existing runs, keep paragraph properties (level, alignment, bullet)
    for r in list(para.runs):
        r._r.getparent().remove(r._r)

    for text, style in parts:
        run = para.add_run()
        run.text = text
        if size is not None:
            run.font.size = size
        if name:
            run.font.name = name
        run.font.bold = style.startswith("b")

        if style.endswith("sub"):
            run.font._rPr.set("baseline", SUB)
        elif style.endswith("sup"):
            run.font._rPr.set("baseline", SUP)


def body_of(slide):
    """First non-title shape that has text."""
    title = slide.shapes.title
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if title is not None and sh._element is title._element:
            continue
        if sh.text_frame.text.strip():
            return sh
    return None


def paras(slide):
    return body_of(slide).text_frame.paragraphs


# ---------------------------------------------------------------- patches
def patch_14(prs):
    p = paras(prs.slides[13])
    rebuild(p[0], [
        ("Real CKKS now replaces the previous placeholder (", ""),
        ("SimulatedCKKSVector", ""),
        ("), so the FHE-vs-plaintext comparison is a controlled experiment "
         "for the first time", ""),
    ])
    # the broken clause: "...quantization in now it is attributable..."
    rebuild(p[1], [
        ("Accuracy is encryption-invariant → any accuracy change under "
         "INT8 quantization is therefore attributable to quantization alone, "
         "not to encryption", ""),
    ])
    rebuild(p[2], [
        ("Measured cryptographic cost per client update (44,164 parameters): "
         "encrypt ~42 ms · encrypted distance ~118–230 ms · "
         "homomorphic aggregation ~5–52 ms · ciphertext on the wire "
         "2.47 MB", ""),
    ])
    print("  slide 14  grammar + wording fixed")


def patch_20(prs):
    p = paras(prs.slides[19])

    rebuild(p[1], [
        ("Each device generates its ", ""),
        ("own secret key ", "b"), ("S", "b"), ("i", "bsub"),
        (" and corresponding ", ""),
        ("public key ", "b"), ("P", "b"), ("i", "bsub"),
    ])

    rebuild(p[2], [
        ("The public keys are combined into an ", ""),
        ("aggregated public key", "b"),
        (":  P̃ = P", ""), ("1", "sub"),
        (" + P", ""), ("2", "sub"),
        (" + ⋯ + P", ""), ("N", "sub"),
    ])

    rebuild(p[3], [
        ("Devices perform ", ""), ("local training", "b"),
        (" and encrypt their model updates under the ", ""),
        ("aggregated public key P̃", "b"),
        (" — not under their own P", ""), ("i", "sub"),
    ])

    rebuild(p[4], [
        ("The server performs ", ""), ("homomorphic addition", "b"),
        (" on the ciphertexts without seeing individual model updates", ""),
    ])

    rebuild(p[5], [
        ("The server obtains an ", ""), ("encrypted aggregate", "b"),
        (" and returns the aggregated ciphertext component to the devices", ""),
    ])

    rebuild(p[6], [
        ("Each device uses its ", ""),
        ("secret key S", "b"), ("i", "bsub"),
        (" and the aggregated ciphertext", "b"),
        (" to generate a ", ""),
        ("decryption share D", "b"), ("i", "bsub"),
        (", then sends it to the server", ""),
    ])
    print("  slide 20  subscripts applied (S_i, P_i, P~, D_i)")


def patch_24(prs):
    p = paras(prs.slides[23])

    rebuild(p[1], [
        ("The server combines the ", ""), ("decryption shares", "b"),
        (" with the aggregated ciphertext to recover ", ""),
        ("only the sum of the model updates", "b"),
    ])

    rebuild(p[4], [
        ("Federated Averaging", "b"), (":  W", ""), ("t+1", "sup"),
        (" = (1/N) ∑", ""), ("i=1", "sub"), ("N", "sup"),
        (" W", ""), ("i", "sub"),
    ])

    rebuild(p[6], [
        ("Individual updates W", ""), ("1", "sub"),
        (", W", ""), ("2", "sub"),
        (", …, W", ""), ("N", "sub"),
        (" remain protected from the server under the paper’s "
         "threat model", ""),
    ])

    rebuild(p[8], [
        ("xMK-CKKS is designed to resist ", ""),
        ("collusion between the server and k < N − 1 participating devices",
         "b"),
    ])
    print("  slide 24  superscript W^(t+1), subscripts W_1..W_N, spacing fixed")


def patch_25(prs):
    slide = prs.slides[24]
    body = body_of(slide)
    tf = body.text_frame
    ps = tf.paragraphs

    L0 = para_font(ps[0])          # top-level bullet style
    L1 = para_font(ps[4])          # sub-bullet style

    rows = [
        (0, [("Ciphertext serialization & comms protocol: ", ""),
             ("partial", "b"),
             (" — authenticated transport channel complete; "
              "serialization + fragmentation spec open", "")]),
        (0, [("Model quantization: ", ""), ("not started", "b"),
             (" — the encryption-neutral accuracy baseline is now "
              "established as its control", "")]),
        (0, [("On-device inference engine / embedded executable: pathway "
              "defined via the OpenFHE + C++ core migration", "")]),
        (0, [("Multi-key CKKS: ", ""), ("xMK-CKKS selected", "b"),
             (" — the homomorphic arithmetic is real CKKS today; the key "
              "management is still an algebraic model", "")]),
        (0, [("Future plan:", "b")]),
        (1, [("OpenFHE migration — makes xMK-CKKS real rather than "
              "modelled", "")]),
        (1, [("Per-round aggregated public key P̃ over the participant "
              "set — removes the dominant dropout case at no extra cost",
              "")]),
        (1, [("Slot-batched distance decryption — keeps a round at 2 "
              "distributed decryptions instead of N + 1", "")]),
        (1, [("Low-magnitude persistent-poisoning experiment — the known "
              "blind spot in the scoring function", "")]),
        (1, [("Scale from 10 to 50 clients; re-acquire raw VeReMi with "
              "sender IDs", "")]),
    ]

    # reuse existing paragraphs, add more if needed
    while len(tf.paragraphs) < len(rows):
        tf.add_paragraph()
    ps = tf.paragraphs

    for i, (level, parts) in enumerate(rows):
        ps[i].level = level
        rebuild(ps[i], parts, base=(L0 if level == 0 else L1))

    # blank any paragraphs left over beyond our content
    for extra in ps[len(rows):]:
        for r in list(extra.runs):
            r._r.getparent().remove(r._r)

    print("  slide 25  rewritten to match the xMK-CKKS narrative")


def main():
    prs = Presentation(str(DECK))
    print("patching %s (%d slides)" % (DECK.name, len(prs.slides)))
    patch_14(prs)
    patch_20(prs)
    patch_24(prs)
    patch_25(prs)
    prs.save(str(DECK))
    print("saved ->", DECK)


if __name__ == "__main__":
    main()
