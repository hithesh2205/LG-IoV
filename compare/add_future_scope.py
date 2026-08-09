from pptx import Presentation
from pptx.util import Inches, Pt
import sys

def add_future_scope(ppt_path, output_path):
    prs = Presentation(ppt_path)
    
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    new_slide = prs.slides.add_slide(slide_layout)
    
    title = new_slide.shapes.title
    title.text = "Future Scope"
    
    content = new_slide.placeholders[1]
    tf = content.text_frame
    
    tf.text = "Hardware Benchmarking: Deploying and testing the ChebyKAN FHE pipeline on real-world Edge and On-Board Unit (OBU) vehicular hardware."
    
    p = tf.add_paragraph()
    p.text = "Genetic Algorithm Optimization: Completing the GA hyperparameter evolution loop to automatically tune the model across diverse non-IID conditions."
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "Latency Reduction: Further optimizing Homomorphic Encryption configurations to push aggregation latency down for massive fleet sizes."
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "Dynamic Vehicle Churn Management: Improving resilience mechanisms for vehicles dropping in and out of the network range during the FHE aggregation phase."
    p.level = 0
    
    prs.save(output_path)
    print(f"Added Future Scope slide and saved to {output_path}")

if __name__ == "__main__":
    add_future_scope("ChebyKAN_updated.pptx", "ChebyKAN_updated.pptx")
