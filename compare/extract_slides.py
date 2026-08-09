from pptx import Presentation
import sys

def extract_text(ppt_path):
    try:
        prs = Presentation(ppt_path)
    except Exception as e:
        return f"Error loading presentation: {e}"
    slides_text = []
    for i, slide in enumerate(prs.slides):
        slide_text = f"--- Slide {i+1} ---\n"
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text += shape.text + "\n"
        slides_text.append(slide_text)
    return "\n".join(slides_text)

if __name__ == "__main__":
    ppt1 = sys.argv[1]
    ppt2 = sys.argv[2]
    
    with open("ppt1_text.txt", "w", encoding="utf-8") as f:
        f.write(extract_text(ppt1))
        
    with open("ppt2_text.txt", "w", encoding="utf-8") as f:
        f.write(extract_text(ppt2))
