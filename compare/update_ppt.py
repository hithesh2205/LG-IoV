from pptx import Presentation
from pptx.util import Inches, Pt
import sys

def add_takeaway_slide(ppt_path, output_path):
    prs = Presentation(ppt_path)
    
    # We want to insert after the title slide (index 0).
    # python-pptx doesn't have a direct 'insert_slide' at index, but we can add one and move it.
    # Actually, we can just add a slide, then move it in the XML, or we can just create it.
    # To move a slide in python-pptx:
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    new_slide = prs.slides.add_slide(slide_layout)
    
    title = new_slide.shapes.title
    title.text = "Takeaways & Discussions from Last Meeting"
    
    content = new_slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Model Evolution: Transitioned from Fourier-based KANConvNet to a Pure Chebyshev Polynomial KAN (ChebyKAN) trunk for optimized parameter size (44,164 parameters)."
    
    p = tf.add_paragraph()
    p.text = "Security Architecture Overhaul: Completely deprecated and removed transport-layer security overhead (AES/ECDH handshakes)."
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "Enhanced Cryptography: Shifted from selective Layer-Wise FHE to Full Layer-wise Homomorphic Encryption (every weight/bias tensor is encrypted separately)."
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "Robust Aggregation: Aggregator now performs Homomorphic FedAvg combined with Multi-Krum outlier rejection directly on FHE ciphertexts."
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "Progress Validation: Successfully ran federated simulation smoke tests across all 4 target datasets to prove the new FHE pipeline."
    p.level = 0
    
    # Move the new slide (which is at the end) to index 1
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[-1])
    xml_slides.insert(1, slides[-1])
    
    prs.save(output_path)
    print(f"Saved updated presentation to {output_path}")

if __name__ == "__main__":
    add_takeaway_slide("ChebyKAN.pptx", "ChebyKAN_updated.pptx")
