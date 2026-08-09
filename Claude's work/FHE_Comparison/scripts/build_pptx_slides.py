"""Append FL-noFHE vs FL-FHE comparison slides + CKKS cost/multi-key slides to a
COPY of the project's deck (never touches the original file).

Reads:
  checkpoints/comparison_summary.json   (from build_comparison.py)
  checkpoints/real_ckks_benchmark.json  (from benchmark_real_ckks.py)

Writes new slides into: reports/ChebyKAN_updated_v2_FHE_comparison.pptx

This is the archival copy of the script (lives in Claude's work/FHE_Comparison/scripts).
Reads its data from this project folder's own results/ directory, not NewModel/checkpoints.
"""
from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

SCRIPTS_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPTS_DIR.parent  # .../Claude's work/FHE_Comparison
DECK_PATH = PROJECT_DIR / "reports" / "ChebyKAN_updated_v2_FHE_comparison.pptx"
CHECKPOINT_DIR = PROJECT_DIR / "results"

# --- Palette (matches the CKKS interactive artifact for cross-deliverable consistency) ---
INK = RGBColor(0x14, 0x1A, 0x20)
INK_DIM = RGBColor(0x57, 0x62, 0x6C)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xEE, 0xF1, 0xF3)
ACCENT = RGBColor(0x1E, 0x8F, 0x80)      # teal — FHE / encrypted
ACCENT_SOFT = RGBColor(0xE1, 0xF0, 0xEE)
WARN = RGBColor(0xA6, 0x6A, 0x15)        # amber — no-FHE / plaintext contrast
WARN_SOFT = RGBColor(0xF6, 0xEC, 0xDA)
LINE = RGBColor(0xD9, 0xE0, 0xE5)
NAVY_DARK = RGBColor(0x0E, 0x2A, 0x2A)

DATASET_LABELS = {
    "can_vtc": "CAN-VTC",
    "car_hack": "Car-Hacking",
    "cicids": "CICIDS-2017",
    "veremi": "VeReMi",
}


def load_json(path: Path):
    if not path.exists():
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def set_title(slide, text, size=30, color=INK):
    tf = slide.shapes.title.text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = color
            r.font.name = "Calibri"


def add_eyebrow(slide, text, top=Inches(0.35)):
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(8), Inches(0.35))
    tf = box.text_frame
    tf.text = text.upper()
    p = tf.paragraphs[0]
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = ACCENT
    p.runs[0].font.name = "Calibri"
    try:
        p.runs[0].font._rPr.set("spc", "150")
    except Exception:
        pass
    return box


def style_table(table, header_fill=ACCENT, header_font=PAPER, body_font=INK, alt_fill=SURFACE):
    for ci, cell in enumerate(table.rows[0].cells):
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(13)
                r.font.color.rgb = header_font
                r.font.name = "Calibri"
    for ri in range(1, len(table.rows)):
        for ci, cell in enumerate(table.rows[ri].cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = alt_fill if ri % 2 == 0 else PAPER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.size = Pt(12.5)
                    r.font.color.rgb = body_font
                    r.font.name = "Calibri"


def build():
    prs = Presentation(str(DECK_PATH))
    layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}

    comparison = load_json(CHECKPOINT_DIR / "comparison_summary.json") or []
    bench = load_json(CHECKPOINT_DIR / "real_ckks_benchmark.json") or []
    bench_by_ds = {b["dataset"]: b for b in bench}

    # ---------------- Slide: Section divider ----------------
    s = prs.slides.add_slide(layouts["Section Header"])
    s.shapes.title.text_frame.text = "FL Without FHE vs. Full Layer-wise FHE"
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = INK
            r.font.bold = True
    body_ph = s.placeholders[1]
    body_ph.text_frame.text = (
        "A controlled, from-scratch comparison: identical ChebyKAN model, identical "
        "Dirichlet client partitions (α=0.3), identical seed — FHE on/off is the only variable."
    )
    for p in body_ph.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(16)
            r.font.color.rgb = INK_DIM

    # ---------------- Slide: Methodology ----------------
    s = prs.slides.add_slide(layouts["Title and Content"])
    set_title(s, "What Changed — and What Didn't")
    body = s.placeholders[1].text_frame
    body.clear()
    lines = [
        ("Only variable: FHE on/off.", True),
        ("Both runs use the same ChebyKAN architecture, same 10 simulated vehicles, same Dirichlet non-IID sharding (α=0.3), same 5 rounds × 2 local epochs, same AdamW config, same random seed (2025), across all 4 datasets.", False),
        ("No-FHE run", True),
        ("Server aggregates plaintext state-dict tensors directly — standard layer-wise FedAvg + Multi-Krum, no encryption wrapper at all.", False),
        ("FHE run", True),
        ("Identical code path, but every layer tensor is wrapped in a CKKS ciphertext object before it leaves a client and before the server aggregates it — matching the “Full Layer-wise FHE” design in this deck's Phase 2.", False),
    ]
    first = True
    for text, is_head in lines:
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        run = p.add_run()
        run.text = text
        run.font.size = Pt(15 if is_head else 13.5)
        run.font.bold = is_head
        run.font.color.rgb = ACCENT if is_head else INK_DIM
        run.font.name = "Calibri"
        p.space_after = Pt(4 if is_head else 10)

    # ---------------- Slide: Results table + chart ----------------
    s = prs.slides.add_slide(layouts["Title Only"])
    set_title(s, "Accuracy Is Preserved — FHE Adds No Measurable Loss")
    add_eyebrow(s, "Fresh training runs · 4 datasets · identical config", top=Inches(1.15))

    if comparison:
        rows = 1 + len(comparison)
        cols = 5
        tbl_shape = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.6), Inches(6.4), Inches(0.5 * rows))
        table = tbl_shape.table
        headers = ["Dataset", "No-FHE Acc.", "FHE Acc.", "Δ Accuracy", "Params"]
        for c, h in enumerate(headers):
            table.cell(0, c).text = h
        for ri, row in enumerate(comparison, start=1):
            table.cell(ri, 0).text = DATASET_LABELS.get(row["dataset"], row["dataset"])
            table.cell(ri, 1).text = f"{row['nofhe_accuracy']*100:.2f}%"
            table.cell(ri, 2).text = f"{row['fhe_accuracy']*100:.2f}%"
            delta = row["accuracy_delta"] * 100
            table.cell(ri, 3).text = f"{delta:+.3f}%"
            table.cell(ri, 4).text = f"{row['parameters']:,}"
        style_table(table)
        for row in table.rows:
            row.height = Inches(0.45)

        # Native grouped bar chart
        chart_data = CategoryChartData()
        chart_data.categories = [DATASET_LABELS.get(r["dataset"], r["dataset"]) for r in comparison]
        chart_data.add_series("No-FHE", [round(r["nofhe_accuracy"] * 100, 2) for r in comparison])
        chart_data.add_series("Full Layer-wise FHE", [round(r["fhe_accuracy"] * 100, 2) for r in comparison])
        gframe = s.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.15), Inches(1.6), Inches(5.7), Inches(4.6), chart_data
        )
        chart = gframe.chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "Test Accuracy: No-FHE vs. FHE (%)"
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.number_format = '0.0"%"'
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(9)
        series_colors = [WARN, ACCENT]
        for i, series in enumerate(plot.series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = series_colors[i]
        chart.value_axis.has_major_gridlines = False
        chart.category_axis.tick_labels.font.size = Pt(10)
        chart.value_axis.tick_labels.font.size = Pt(10)

    caption = s.shapes.add_textbox(Inches(0.5), Inches(1.6) + Inches(0.5 * (1 + len(comparison))) + Inches(0.2), Inches(6.4), Inches(1.2))
    caption.text_frame.word_wrap = True
    caption.text_frame.text = (
        "The simulated CKKS layer used for training (fhe_adapter.SimulatedCKKSVector) injects only "
        "gaussian noise of σ=1e-9 to mirror encryption rounding — it validates that the aggregation "
        "math is zero-loss, but it does not model real cryptographic cost. See next slide for that."
    )
    for p in caption.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11)
            r.font.italic = True
            r.font.color.rgb = INK_DIM

    # ---------------- Slide: Real CKKS overhead ----------------
    s = prs.slides.add_slide(layouts["Title Only"])
    set_title(s, "What Real CKKS Actually Costs")
    add_eyebrow(s, "Benchmarked with TenSEAL · real lattice cryptography, not a simulation", top=Inches(1.15))

    bench_rows = list(bench_by_ds.values())
    if bench_rows:
        rows = 1 + len(bench_rows)
        cols = 5
        tbl_shape = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.65), Inches(9.0), Inches(0.5 * rows))
        table = tbl_shape.table
        headers = ["Dataset", "Encrypt/client", "Homomorphic sum", "Decrypt", "Ciphertext expansion"]
        for c, h in enumerate(headers):
            table.cell(0, c).text = h
        for ri, row in enumerate(bench_rows, start=1):
            table.cell(ri, 0).text = DATASET_LABELS.get(row["dataset"], row["dataset"])
            table.cell(ri, 1).text = f"{row['encrypt_time_sec_per_client']*1000:.1f} ms"
            table.cell(ri, 2).text = f"{row['homomorphic_aggregate_time_sec']*1000:.1f} ms"
            table.cell(ri, 3).text = f"{row['decrypt_time_sec']*1000:.1f} ms"
            table.cell(ri, 4).text = f"{row['ciphertext_expansion_factor']:.1f}×"
        style_table(table)
        for row in table.rows:
            row.height = Inches(0.45)

    note = s.shapes.add_textbox(Inches(0.5), Inches(1.65) + Inches(0.5 * (1 + len(bench_rows))) + Inches(0.35), Inches(9.0), Inches(2.2))
    note.text_frame.word_wrap = True
    note.text_frame.text = "Method"
    note.text_frame.paragraphs[0].runs[0].font.bold = True
    note.text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT
    note.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    p2 = note.text_frame.add_paragraph()
    p2.text = (
        "Real TenSEAL CKKS (poly_modulus_degree=8192, coeff_mod=[60,40,40,60]) encrypting the actual "
        "ChebyKAN parameter tensors, aggregating 7 active clients (of 10, sampling rate 0.7) homomorphically "
        "per round — the true crypto cost that the numpy simulation on the previous slide doesn't capture. "
        "Numerical error stayed below ~3×10⁻⁸ versus plaintext FedAvg."
    )
    p2.runs[0].font.size = Pt(13)
    p2.runs[0].font.color.rgb = INK_DIM

    # ---------------- Slide: Multi-key / Threshold CKKS ----------------
    s = prs.slides.add_slide(layouts["Two Content"])
    set_title(s, "Beyond a Single Shared Key")
    left = s.placeholders[1].text_frame
    left.clear()
    left.word_wrap = True
    left_lines = [
        ("The single-key problem", True),
        ("Every client in this deck's design (and in TenSEAL/SEAL) shares one secret key. Any one compromised vehicle can decrypt every other vehicle's updates — a single point of failure across the whole fleet.", False),
        ("Threshold / multiparty CKKS (recommended)", True),
        ("One shared public key; the secret key is split into per-vehicle shares. Decryption needs a quorum of partial-decryption shares combined — no single vehicle ever holds enough to decrypt alone. Same ciphertext math as single-key CKKS, so it's a direct extension of this project's existing pipeline.", False),
    ]
    first = True
    for text, is_head in left_lines:
        p = left.paragraphs[0] if first else left.add_paragraph()
        first = False
        r = p.add_run()
        r.text = text
        r.font.size = Pt(14 if is_head else 12.5)
        r.font.bold = is_head
        r.font.color.rgb = ACCENT if is_head else INK_DIM
        r.font.name = "Calibri"
        p.space_after = Pt(4 if is_head else 12)

    right = s.placeholders[2].text_frame
    right.clear()
    right.word_wrap = True
    right_lines = [
        ('"True" multi-key CKKS', True),
        ("Every party keeps a fully independent key; ciphertexts from different keys are algebraically extended before they can be combined, with its own key-switching and noise-growth analysis. Active research area — not supported out of the box by TenSEAL/SEAL.", False),
        ("Recommended next step", True),
        ("Prototype the threshold variant with a library built for it (e.g. Lattigo's dckks / mkckks packages) rather than extending TenSEAL, which is single-key by design.", False),
    ]
    first = True
    for text, is_head in right_lines:
        p = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        r = p.add_run()
        r.text = text
        r.font.size = Pt(14 if is_head else 12.5)
        r.font.bold = is_head
        r.font.color.rgb = WARN if is_head else INK_DIM
        r.font.name = "Calibri"
        p.space_after = Pt(4 if is_head else 12)

    prs.save(str(DECK_PATH))
    _fix_jpg_content_type(DECK_PATH)
    print(f"[build_pptx_slides] Saved {DECK_PATH} with {len(prs.slides._sldIdLst)} slides")


def _fix_jpg_content_type(deck_path: Path) -> None:
    """python-pptx sometimes re-serializes a pre-existing jpg image part with the
    non-standard MIME type 'image/jpg' instead of the correct 'image/jpeg'. Patch
    [Content_Types].xml in place so the saved package stays spec-compliant.
    """
    tmp_path = deck_path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(deck_path, "r") as zin:
        names = zin.namelist()
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names:
                data = zin.read(name)
                if name == "[Content_Types].xml":
                    data = data.replace(b'ContentType="image/jpg"', b'ContentType="image/jpeg"')
                    if b'Extension="jpg"' not in data:
                        data = data.replace(
                            b'<Default Extension="jpeg" ContentType="image/jpeg"/>',
                            b'<Default Extension="jpeg" ContentType="image/jpeg"/>'
                            b'<Default Extension="jpg" ContentType="image/jpeg"/>',
                        )
                zout.writestr(name, data)
    shutil.move(str(tmp_path), str(deck_path))


if __name__ == "__main__":
    build()
